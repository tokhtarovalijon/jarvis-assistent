"""
Генерация презентаций (.pptx): структура через ИИ, реальные экономические данные World Bank,
фото-фон с Pexels, фирменный тёмный дизайн, заметки докладчика и глоссарий терминов.
"""
import os
import re
import random
import requests
from datetime import datetime
from pptx import Presentation
from pptx.util import Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from config import client, PEXELS_API_KEY, BOT_NAME
from tts import speak

# --- ГЕНЕРАЦИЯ ПРЕЗЕНТАЦИЙ (.pptx) ---
PRESENTATIONS_DIR = "презентации"
os.makedirs(PRESENTATIONS_DIR, exist_ok=True)

PRESENTATION_TRIGGERS = ["создай презентацию", "сделай презентацию", "презентация про", "презентация на тему"]

# Фирменные цвета презентации (в стиле HUD-интерфейса Джарвиса)
PPTX_BG_DARK = RGBColor(0x0A, 0x14, 0x24)
PPTX_ACCENT_ORANGE = RGBColor(0xFF, 0x8C, 0x00)
PPTX_TEXT_LIGHT = RGBColor(0xE8, 0xF0, 0xFF)
PPTX_TEXT_MUTED = RGBColor(0x8A, 0xA5, 0xC2)


def is_presentation_request(text):
    return any(trigger in text for trigger in PRESENTATION_TRIGGERS)


# --- РЕАЛЬНЫЕ ЭКОНОМИЧЕСКИЕ ДАННЫЕ (World Bank Open Data — бесплатно, без ключа, любая страна) ---
WORLD_BANK_INDICATORS = {
    "ВВП (текущие цены, $)": "NY.GDP.MKTP.CD",
    "Рост ВВП (%)": "NY.GDP.MKTP.KD.ZG",
    "Инфляция (%)": "FP.CPI.TOTL.ZG",
    "Население": "SP.POP.TOTL",
    "ВВП на душу населения ($)": "NY.GDP.PCAP.CD",
    "Безработица (%)": "SL.UEM.TOTL.ZS",
}


def detect_country_code(topic):
    """Просит ИИ определить ISO-2 код страны, если тема касается конкретной страны."""
    if not client:
        return None
    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Определи, упоминается ли в теме конкретная страна. Если да — ответь "
                        "ТОЛЬКО её ISO 3166-1 alpha-2 кодом (например, TJ для Таджикистана, "
                        "RU для России, US для США). Если страна не упоминается — ответь "
                        "ровно словом НЕТ, без пояснений."
                    ),
                },
                {"role": "user", "content": topic},
            ],
            temperature=0,
        )
        answer = response.choices[0].message.content.strip().upper()
        if len(answer) == 2 and answer.isalpha():
            return answer
        return None
    except Exception as e:
        print(f"[Ошибка определения страны] {e}")
        return None


def fetch_country_economic_data(country_code):
    """Скачивает реальные экономические показатели страны с World Bank Open Data."""
    results = {}
    for label, indicator in WORLD_BANK_INDICATORS.items():
        try:
            url = f"https://api.worldbank.org/v2/country/{country_code}/indicator/{indicator}"
            response = requests.get(url, params={"format": "json", "per_page": 10}, timeout=10)
            response.raise_for_status()
            data = response.json()
            if len(data) < 2 or not data[1]:
                continue
            for entry in data[1]:
                if entry.get("value") is not None:
                    results[label] = f"{entry['value']:,.1f} ({entry['date']} г.)"
                    break
        except Exception as e:
            print(f"[Ошибка World Bank API для {indicator}] {e}")
            continue
    return results


def format_economic_data(data):
    if not data:
        return ""
    lines = ["РЕАЛЬНЫЕ ДАННЫЕ ИЗ WORLD BANK OPEN DATA (используй именно эти цифры):"]
    for label, value in data.items():
        lines.append(f"- {label}: {value}")
    lines.append(
        "\nОБЯЗАТЕЛЬНО используй эти цифры в презентации там, где уместно (например, на слайде "
        "со статистикой). Не заменяй их выдуманными значениями и не противоречь им."
    )
    return "\n".join(lines)


def generate_presentation_outline(topic):
    """Просит ИИ вернуть структуру презентации в строгом текстовом формате, с упором на факты,
    заметками докладчика к каждому слайду и глоссарием терминов в конце."""
    system_prompt = (
        "Ты составляешь структуру презентации для студента экономического университета. "
        "Ответь СТРОГО в этом формате, без лишнего текста:\n\n"
        "ЗАГОЛОВОК: <название презентации>\n"
        "СЛАЙД: <заголовок слайда 1>\n"
        "- <тезис 1>\n"
        "- <тезис 2>\n"
        "- <тезис 3>\n"
        "- <тезис 4>\n"
        "- <тезис 5>\n"
        "ЗАМЕТКИ: <3-5 предложений — что докладчик говорит вслух по этому слайду, объясняя "
        "тезисы простыми словами, как одногруппнику>\n"
        "СЛАЙД: <заголовок слайда 2>\n"
        "...\n"
        "ГЛОССАРИЙ:\n"
        "ТЕРМИН: <термин 1> — <короткое понятное определение>\n"
        "ТЕРМИН: <термин 2> — <короткое понятное определение>\n"
        "...\n\n"
        "Требования:\n"
        "- Сделай 10-12 слайдов, раскрывающих тему глубоко и по существу.\n"
        "- В каждом слайде 5-6 тезисов (не меньше). КАЖДЫЙ тезис — законченное объясняющее "
        "предложение с конкретным фактом (цифра, дата, процент, закон/организация/событие) "
        "И пояснением, что это значит или почему это важно. НЕ пиши тезисы в формате голой "
        "пары 'Показатель: значение (год)' — это неинформативно. Вместо 'ВВП: 17.7 млрд "
        "(2025 г.)' пиши 'ВВП вырос до 17.7 млрд долларов в 2025 году — рост на 8.4% отражает "
        "восстановление экономики после спада'.\n"
        "- Тезисы длиной 15-25 слов — достаточно, чтобы дать контекст, но не превращать слайд "
        "в сплошной текст.\n"
        "- К КАЖДОМУ слайду обязательно добавь строку ЗАМЕТКИ — развёрнутое объяснение на "
        "3-5 предложений для докладчика, а не повтор тезисов другими словами.\n"
        "- В конце добавь ГЛОССАРИЙ из 6-10 ключевых терминов темы с короткими определениями.\n"
        "- Один из слайдов должен содержать конкретную статистику или динамику показателей — "
        "и она тоже должна быть подана как объясняющие предложения, а не список 'ярлык: цифра'.\n"
        "- Будь точен с датами и цифрами: если не уверен — не выдумывай, используй общепринятые "
        "и хорошо известные данные. Если ниже даны реальные данные — используй именно их, но "
        "оформи их как содержательные предложения с интерпретацией.\n"
        "- НЕ используй markdown-разметку (никаких **, ##, > и т.п.) — только обычный текст, "
        "строго в указанном формате, без вступлений и заключений."
    )

    PRESENTATION_ANGLES = [
        "с упором на исторический контекст и то, как тема развивалась со временем",
        "с упором на свежую статистику и динамику показателей последних лет",
        "с упором на сравнение с другими странами или примерами для контраста",
        "с упором на причины проблемы и её практические последствия",
        "с упором на разные точки зрения и школы экономической мысли по теме",
        "с упором на практические примеры и кейсы из реальной жизни",
        "с упором на прогнозы и возможные сценарии развития темы в будущем",
        "с упором на ключевые цифры, факты и их интерпретацию",
    ]
    chosen_angle = random.choice(PRESENTATION_ANGLES)

    user_content = (
        f"Тема презентации: {topic}\n\n"
        f"Раскрой эту тему {chosen_angle}. Это одна из множества презентаций на похожие темы — "
        "выбери свою структуру и порядок слайдов самостоятельно, не подгоняй под шаблон."
    )

    country_code = detect_country_code(topic)
    if country_code:
        econ_data = fetch_country_economic_data(country_code)
        grounding = format_economic_data(econ_data)
        if grounding:
            user_content += f"\n\n{grounding}"

    def call_model(reasoning_effort="low", tokens=6000):
        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_content},
            ],
            temperature=0.65,
            max_tokens=tokens,
            reasoning_effort=reasoning_effort,
        )
        return (response.choices[0].message.content or "").strip()

    try:
        result = call_model("low", 8000)
    except Exception as e:
        print(f"[Groq ошибка, параметр reasoning_effort не поддержан?] {e}")
        result = ""

    if not result:
        # Пустой ответ — модель, вероятно, потратила лимит на "размышления".
        # Пробуем ещё раз без reasoning_effort и с более настойчивым промптом.
        print("[Отладка] Пустой ответ от модели, повторяю запрос...")
        try:
            response = client.chat.completions.create(
                model="openai/gpt-oss-20b",
                messages=[
                    {"role": "system", "content": system_prompt},
                    {
                        "role": "user",
                        "content": user_content + "\n\nВАЖНО: ответь СРАЗУ структурой, без внутренних рассуждений.",
                    },
                ],
                temperature=0.65,
                max_tokens=10000,
            )
            result = (response.choices[0].message.content or "").strip()
        except Exception as e:
            print(f"[Groq ошибка при повторной попытке] {e}")
            result = ""

    return result


def parse_outline(raw_text):
    """Разбирает текстовый ответ ИИ на структуру: заголовок, слайды (с тезисами и заметками
    докладчика) и глоссарий терминов. Устойчив к markdown-разметке (**, #, > и т.п.)."""
    title = "Презентация"
    slides = []
    glossary = []
    current_slide = None
    mode = None  # None | "notes" | "glossary"

    for raw_line in raw_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        # Убираем разметку markdown в начале и конце строки
        cleaned = re.sub(r'^[#>\*\s\d\.\)]+', '', line).strip()
        cleaned = cleaned.rstrip('*').strip()
        upper_cleaned = cleaned.upper()

        if upper_cleaned.startswith("ЗАГОЛОВОК"):
            parts = cleaned.split(":", 1)
            if len(parts) > 1:
                title = parts[1].strip().strip('*').strip()
            mode = None

        elif upper_cleaned.startswith("СЛАЙД"):
            if current_slide:
                slides.append(current_slide)
            parts = cleaned.split(":", 1)
            slide_title = parts[1].strip().strip('*').strip() if len(parts) > 1 else "Слайд"
            current_slide = {"title": slide_title, "bullets": [], "notes": ""}
            mode = None

        elif upper_cleaned.startswith("ЗАМЕТКИ") and current_slide is not None:
            parts = cleaned.split(":", 1)
            current_slide["notes"] = parts[1].strip() if len(parts) > 1 else ""
            mode = "notes"

        elif upper_cleaned.startswith("ГЛОССАРИЙ"):
            if current_slide:
                slides.append(current_slide)
                current_slide = None
            mode = "glossary"

        elif upper_cleaned.startswith("ТЕРМИН") and mode == "glossary":
            parts = cleaned.split(":", 1)
            if len(parts) > 1:
                glossary.append(parts[1].strip())

        elif (line.lstrip().startswith("-") or line.lstrip().startswith("•")) and current_slide and mode != "glossary":
            bullet = cleaned.lstrip("-•").strip()
            if bullet:
                current_slide["bullets"].append(bullet)
            mode = None

        elif mode == "notes" and current_slide is not None:
            # Продолжение заметок докладчика на следующей строке (если ИИ перенёс текст)
            current_slide["notes"] = (current_slide["notes"] + " " + cleaned).strip()

    if current_slide:
        slides.append(current_slide)

    if not slides:
        print("[Отладка] Не удалось распознать структуру. Сырой ответ ИИ:")
        print(raw_text)

    return title, slides, glossary


def translate_topic_for_search(topic):
    """Переводит тему в 2-3 английских ключевых слова для поиска фото на Pexels."""
    if not client:
        return topic
    try:
        response = client.chat.completions.create(
            model="openai/gpt-oss-20b",
            messages=[
                {
                    "role": "system",
                    "content": (
                        "Переведи тему презентации в 2-3 английских ключевых слова для поиска "
                        "фотографии по смыслу темы (например, экономика/бизнес/финансы визуально). "
                        "Ответь ТОЛЬКО ключевыми словами на английском, без пунктуации и пояснений."
                    ),
                },
                {"role": "user", "content": topic},
            ],
            temperature=0.2,
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        print(f"[Ошибка перевода темы] {e}")
        return topic


def fetch_pexels_image(query):
    """Скачивает фото по теме с Pexels. Возвращает путь к файлу или None."""
    if not PEXELS_API_KEY:
        return None
    try:
        response = requests.get(
            "https://api.pexels.com/v1/search",
            headers={"Authorization": PEXELS_API_KEY},
            params={"query": query, "per_page": 1, "orientation": "landscape"},
            timeout=10,
        )
        response.raise_for_status()
        data = response.json()
        photos = data.get("photos", [])
        if not photos:
            return None
        image_url = photos[0]["src"]["large2x"]
        image_response = requests.get(image_url, timeout=15)
        image_response.raise_for_status()
        image_path = os.path.join(PRESENTATIONS_DIR, "_bg_temp.jpg")
        with open(image_path, "wb") as f:
            f.write(image_response.content)
        return image_path
    except Exception as e:
        print(f"[Ошибка загрузки фото Pexels] {e}")
        return None


def build_pptx(title, slides, image_path=None, glossary=None):
    prs = Presentation()
    slide_width = prs.slide_width
    slide_height = prs.slide_height

    # --- ТИТУЛЬНЫЙ СЛАЙД ---
    blank_layout = prs.slide_layouts[6]
    title_slide = prs.slides.add_slide(blank_layout)

    if image_path and os.path.exists(image_path):
        title_slide.shapes.add_picture(image_path, 0, 0, width=slide_width, height=slide_height)

        # Тёмная полоса-подложка снизу для читаемости текста
        band_height = int(slide_height * 0.32)
        band = title_slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, slide_height - band_height, slide_width, band_height
        )
        band.fill.solid()
        band.fill.fore_color.rgb = RGBColor(0x00, 0x00, 0x00)
        band.line.fill.background()

        title_box = title_slide.shapes.add_textbox(
            Inches(0.5), slide_height - band_height + Inches(0.15), slide_width - Inches(1), Inches(1.2)
        )
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        subtitle_box = title_slide.shapes.add_textbox(
            Inches(0.5), slide_height - Inches(0.55), slide_width - Inches(1), Inches(0.4)
        )
        sp = subtitle_box.text_frame.paragraphs[0]
        sp.text = f"Подготовлено ассистентом {BOT_NAME}"
        sp.font.size = Pt(14)
        sp.font.color.rgb = PPTX_ACCENT_ORANGE
    else:
        # Фото не получилось — фирменный тёмный титульный слайд без фото
        title_slide.background.fill.solid()
        title_slide.background.fill.fore_color.rgb = PPTX_BG_DARK

        title_box = title_slide.shapes.add_textbox(
            Inches(0.7), slide_height / 2 - Inches(0.8), slide_width - Inches(1.4), Inches(1.5)
        )
        tf = title_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = PPTX_ACCENT_ORANGE

        subtitle_box = title_slide.shapes.add_textbox(
            Inches(0.7), slide_height / 2 + Inches(0.5), slide_width - Inches(1.4), Inches(0.5)
        )
        sp = subtitle_box.text_frame.paragraphs[0]
        sp.text = f"Подготовлено ассистентом {BOT_NAME}"
        sp.font.size = Pt(16)
        sp.font.color.rgb = PPTX_TEXT_MUTED

    # --- СЛАЙДЫ С КОНТЕНТОМ (фирменный тёмный дизайн) ---
    blank_layout = prs.slide_layouts[6]
    for idx, s in enumerate(slides, start=1):
        slide = prs.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PPTX_BG_DARK

        # Оранжевая акцентная полоса сверху
        accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_width, Inches(0.12))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = PPTX_ACCENT_ORANGE
        accent_bar.line.fill.background()

        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), slide_width - Inches(1.2), Inches(0.9))
        tp = title_box.text_frame.paragraphs[0]
        tp.text = s["title"]
        tp.font.size = Pt(28)
        tp.font.bold = True
        tp.font.color.rgb = PPTX_ACCENT_ORANGE

        body_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.3), slide_width - Inches(1.4), Inches(5.6))
        body = body_box.text_frame
        body.word_wrap = True
        for i, bullet in enumerate(s["bullets"]):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = f"•  {bullet}"
            p.font.size = Pt(15)
            p.font.color.rgb = PPTX_TEXT_LIGHT
            p.space_after = Pt(8)

        # Номер слайда
        page_box = slide.shapes.add_textbox(
            slide_width - Inches(1), slide_height - Inches(0.5), Inches(0.7), Inches(0.3)
        )
        page_p = page_box.text_frame.paragraphs[0]
        page_p.text = str(idx)
        page_p.font.size = Pt(12)
        page_p.font.color.rgb = PPTX_TEXT_MUTED

        # Заметки докладчика — видны только в режиме показа заметок, не мешают слайду
        notes = s.get("notes", "").strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    # --- СЛАЙД-ГЛОССАРИЙ (ключевые термины темы) ---
    if glossary:
        glossary_slide = prs.slides.add_slide(blank_layout)
        glossary_slide.background.fill.solid()
        glossary_slide.background.fill.fore_color.rgb = PPTX_BG_DARK

        accent_bar = glossary_slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, slide_width, Inches(0.12))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = PPTX_ACCENT_ORANGE
        accent_bar.line.fill.background()

        title_box = glossary_slide.shapes.add_textbox(
            Inches(0.6), Inches(0.35), slide_width - Inches(1.2), Inches(0.9)
        )
        tp = title_box.text_frame.paragraphs[0]
        tp.text = "Глоссарий терминов"
        tp.font.size = Pt(28)
        tp.font.bold = True
        tp.font.color.rgb = PPTX_ACCENT_ORANGE

        body_box = glossary_slide.shapes.add_textbox(
            Inches(0.7), Inches(1.4), slide_width - Inches(1.4), Inches(5.2)
        )
        body = body_box.text_frame
        body.word_wrap = True
        for i, term_def in enumerate(glossary):
            p = body.paragraphs[0] if i == 0 else body.add_paragraph()
            p.text = f"•  {term_def}"
            p.font.size = Pt(16)
            p.font.color.rgb = PPTX_TEXT_LIGHT
            p.space_after = Pt(8)

    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    safe_title = re.sub(r'[\\/*?:"<>|]', "", title)[:50]
    filename = os.path.join(PRESENTATIONS_DIR, f"{safe_title}_{timestamp}.pptx")
    prs.save(filename)

    if image_path and os.path.exists(image_path):
        try:
            os.remove(image_path)
        except Exception:
            pass

    return filename


def create_presentation(topic):
    if not client:
        speak("ИИ-модуль не настроен, не могу создать презентацию, сэр.")
        return
    speak(f"Готовлю презентацию на тему: {topic}. Это займёт немного времени, сэр.")
    try:
        raw_outline = generate_presentation_outline(topic)
        title, slides, glossary = parse_outline(raw_outline)
        if not slides:
            speak("Не удалось разобрать структуру презентации, сэр. Попробуйте переформулировать тему.")
            return

        image_path = None
        if PEXELS_API_KEY:
            search_query = translate_topic_for_search(topic)
            image_path = fetch_pexels_image(search_query)

        filepath = build_pptx(title, slides, image_path, glossary)
        speak("Презентация готова, сэр. Сохранена и открывается.")
        os.startfile(os.path.abspath(filepath))
    except Exception as e:
        print(f"[Ошибка создания презентации] {e}")
        speak("Произошла ошибка при создании презентации, сэр.")

